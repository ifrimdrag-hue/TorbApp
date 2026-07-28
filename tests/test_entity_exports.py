"""Exports for the entity detail pages: context, Excel round-trip, PPT, authz."""

import io

import pytest
from pptx import Presentation

AN = 2026
AGENT = 'Agent Test'
CLIENT = 'C001'
BRAND = 'Basilur'
SKU = 'SKU001'
EMPTY_MONTH = 2  # the seed only has month 1, so month 2 yields no rows


def test_agent_kpi_exposes_marja_neta_pct(flask_app):
    import queries
    with flask_app.app_context():
        kpi = queries.agent_kpi(AGENT, AN)
    assert 'marja_neta_pct' in kpi
    # Seed (tests/conftest.py): marja_bruta 200+150=350, val_neta 500+400=900,
    # no conditii_comerciale rows -> condition cost 0 -> marja_neta == marja_bruta.
    assert kpi['marja_neta'] == 350
    assert kpi['marja_neta_pct'] == pytest.approx(38.9, abs=0.1)


def test_agent_kpi_marja_neta_pct_is_zero_without_sales(flask_app):
    import queries
    with flask_app.app_context():
        kpi = queries.agent_kpi('Agent Inexistent', AN)
    assert kpi['marja_neta_pct'] == 0


# ── Context builders ─────────────────────────────────────────────────────────

ENTITY_IDENT = {
    'client': CLIENT,
    'agent': AGENT,
    'brand': BRAND,
    'produs': SKU,
}

ENTITY_NAV = {
    'client': 'clients',
    'agent': 'team',
    'brand': 'products',
    'produs': 'products',
}


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs'])
def test_build_context_shape(flask_app, entity):
    from exports import entities
    with flask_app.app_context():
        ctx = entities.build_context(entity, ENTITY_IDENT[entity], AN, None)
    assert ctx is not None
    assert ctx['nav'] == ENTITY_NAV[entity]
    assert ctx['title']
    assert ctx['filename_base']
    assert 1 <= len(ctx['cards']) <= 4
    assert all(isinstance(v, str) for _, v in ctx['cards'])
    assert len(ctx['sheets']) == 4
    for table in ctx['tables']:
        assert set(table) >= {'title', 'headers', 'rows', 'left', 'width'}


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs'])
def test_build_context_unknown_ident_returns_none(flask_app, entity):
    from exports import entities
    with flask_app.app_context():
        assert entities.build_context(entity, 'NU_EXISTA_XYZ', AN, None) is None


def test_build_context_unknown_entity_returns_none(flask_app):
    from exports import entities
    with flask_app.app_context():
        assert entities.build_context('nope', 'x', AN, None) is None


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs'])
def test_build_context_honours_luna(flask_app, entity):
    """Month 2 has no seeded rows, so its period sheets must come back empty
    while the unfiltered context has data.

    client's resolve query (client_info) is period-independent, so an empty
    month still resolves and must show empty period sheets. agent/brand/produs
    resolve on a period-scoped KPI query, so an empty month makes the builder
    return None outright — that None IS the period filter working, and is
    asserted explicitly rather than folded into an `is None or ...` disjunction
    that would let a broken luna filter on the *list* queries pass unnoticed.
    """
    from exports import entities
    with flask_app.app_context():
        full = entities.build_context(entity, ENTITY_IDENT[entity], AN, None)
        filtered = entities.build_context(entity, ENTITY_IDENT[entity], AN, EMPTY_MONTH)
    period_sheet = f'Clienți {AN}' if entity != 'client' else f'Produse {AN}'
    assert full['sheets'][period_sheet]
    if entity == 'client':
        assert filtered is not None
        assert not filtered['sheets'][period_sheet]
        assert not filtered['sheets']['Brand Mix']
    else:
        assert filtered is None


def test_period_label(flask_app):
    from exports import entities
    with flask_app.app_context():
        assert entities._period_label(2026, 7, None) == '2026 · Iulie'
        assert entities._period_label(2026, None, 7) == '2026 · Ian–Iul'
        assert entities._period_label(2026, None, 12) == '2026'
        assert entities._period_label(2026, None, None) == '2026'


def test_table_caption_reports_truncation(flask_app):
    from exports import entities
    rows = [{'x': i} for i in range(40)]
    table = entities._table('Clienți', 0.3, 6.9, rows, lambda r: {'X': str(r['x'])}, limit=15)
    assert table['title'] == 'Clienți — top 15 din 40'
    assert len(table['rows']) == 15
    assert table['headers'] == ['X']


def test_produs_context_aggregates_sku_variants(flask_app, monkeypatch):
    """sku_variants expansion must survive — the same article appears in
    tranzactii under several spellings and the page aggregates over all."""
    import queries
    from exports import entities
    seen = {}

    real_kpi = queries.product_kpi

    def spy(sku, an, **kw):
        seen['sku'] = sku
        return real_kpi(sku, an, **kw)

    monkeypatch.setattr(queries, 'product_kpi', spy)
    monkeypatch.setattr(queries, 'sku_variants', lambda s: [s, s + ' VARIANTA'])
    with flask_app.app_context():
        entities.build_context('produs', SKU, AN, None)
    assert seen['sku'] == [SKU, SKU + ' VARIANTA']


# ── Deck builder ─────────────────────────────────────────────────────────────

EXPECTED_SLIDES = {'client': 3, 'agent': 4, 'brand': 3, 'produs': 3}


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs'])
def test_build_entity_ppt_slide_count(flask_app, entity):
    from exports import entities, ppt_export
    with flask_app.app_context():
        ctx = entities.build_context(entity, ENTITY_IDENT[entity], AN, None)
        buf = ppt_export.build_entity_ppt(ctx)
    prs = Presentation(buf)
    assert len(prs.slides) == EXPECTED_SLIDES[entity]


def test_build_entity_ppt_cover_shows_period(flask_app):
    from exports import entities, ppt_export
    with flask_app.app_context():
        ctx = entities.build_context('brand', BRAND, AN, 1)
        buf = ppt_export.build_entity_ppt(ctx)
    prs = Presentation(buf)
    texts = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert '2026 · Ianuarie' in texts
    assert 'Brand: Basilur' in texts


def test_build_entity_ppt_without_trend_drops_the_chart_slide(flask_app):
    from exports import entities, ppt_export
    with flask_app.app_context():
        ctx = entities.build_context('brand', BRAND, AN, None)
    ctx['trend'] = {}
    buf = ppt_export.build_entity_ppt(ctx)
    assert len(Presentation(buf).slides) == 2


def test_old_named_builders_are_gone():
    from exports import ppt_export
    assert not hasattr(ppt_export, 'build_client_ppt')
    assert not hasattr(ppt_export, 'build_agent_ppt')
    assert not hasattr(ppt_export, '_slide_client_detail')
    assert not hasattr(ppt_export, '_slide_agent_detail')


# ── PPT route ────────────────────────────────────────────────────────────────

PPT_MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'

PPT_QUERY = {
    'client': f'cod_client={CLIENT}',
    'agent': 'name=Agent+Test',
    'brand': f'furnizor={BRAND}',
    'produs': f'sku={SKU}',
}


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs'])
def test_ppt_route_returns_openable_deck(client, entity):
    rv = client.get(f'/export/ppt/{entity}?{PPT_QUERY[entity]}&an={AN}')
    assert rv.status_code == 200
    assert rv.mimetype == PPT_MIME
    assert rv.data
    prs = Presentation(io.BytesIO(rv.data))
    assert len(prs.slides) == EXPECTED_SLIDES[entity]


@pytest.mark.parametrize('entity', ['dashboard', 'profitabilitate'])
def test_ppt_route_still_serves_overview_decks(client, entity):
    rv = client.get(f'/export/ppt/{entity}?an={AN}')
    assert rv.status_code == 200
    assert rv.mimetype == PPT_MIME
    assert len(Presentation(io.BytesIO(rv.data)).slides) >= 2


def test_ppt_route_unknown_entity_is_404(client):
    assert client.get('/export/ppt/nope').status_code == 404


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs'])
def test_ppt_route_missing_ident_is_404(client, entity):
    assert client.get(f'/export/ppt/{entity}?an={AN}').status_code == 404


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs'])
def test_ppt_route_unknown_ident_is_404(client, entity):
    param = PPT_QUERY[entity].split('=')[0]
    rv = client.get(f'/export/ppt/{entity}?{param}=NU_EXISTA_XYZ&an={AN}')
    assert rv.status_code == 404


@pytest.mark.parametrize('entity', ['client', 'agent', 'brand', 'produs', 'dashboard'])
def test_ppt_route_denies_role_without_nav_access(client, monkeypatch, entity):
    import authz
    monkeypatch.setattr(authz, 'can_access_nav', lambda role, key: False)
    query = PPT_QUERY.get(entity, '')
    rv = client.get(f'/export/ppt/{entity}?{query}&an={AN}')
    assert rv.status_code == 403


def test_old_ppt_endpoints_are_gone(flask_app):
    endpoints = {r.endpoint for r in flask_app.url_map.iter_rules()}
    assert 'reports.export_ppt' in endpoints
    for old in ('reports.export_ppt_client', 'reports.export_ppt_agent',
                'reports.export_ppt_dashboard', 'reports.export_ppt_profitabilitate'):
        assert old not in endpoints


def test_ppt_route_is_allowlisted_not_orphaned():
    import nav_registry as nr
    assert 'reports.export_ppt' in nr.UNGATED_ENDPOINTS
    registered = {ep for item in nr.NAV_REGISTRY for ep in item.endpoints}
    assert 'reports.export_ppt_client' not in registered
    assert 'reports.export_ppt_agent' not in registered
