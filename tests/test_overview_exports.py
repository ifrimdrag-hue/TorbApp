"""Exports for the list/report pages: context, filters, period, Excel, PPT, authz."""

import io

import openpyxl
import pytest
from pptx import Presentation

AN = 2026
EMPTY_MONTH = 2  # the seed only has month 1, so month 2 yields no rows

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PPT_MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'


def _rows(sheet):
    """Sheets are either a plain row list or {'rows', 'headers'} (basilur)."""
    return sheet['rows'] if isinstance(sheet, dict) else sheet


def test_team_context_shape(flask_app):
    """Structure only — the session DB is shared and several other test modules
    insert transactions, so figures are asserted against a monkeypatched query
    in test_team_cards_render_the_period_totals instead."""
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('team', AN, None, {})
    assert ctx['nav'] == 'team'
    assert ctx['filename_base'] == f'echipa_{AN}'
    assert list(ctx['sheets']) == [f'Echipa {AN}', f'Echipa {AN - 1}']
    assert 'Agent Test' in [r['agent'] for r in ctx['sheets'][f'Echipa {AN}']]
    assert len(ctx['cards']) == 4
    assert all(isinstance(v, str) for _, v in ctx['cards'])


def test_team_cards_render_the_period_totals(flask_app, monkeypatch):
    import queries
    from exports import overviews
    monkeypatch.setattr(queries, 'team_table', lambda an, **kw: [
        {'agent': 'A', 'val_neta': 500, 'marja_bruta': 200, 'marja_neta': 200},
        {'agent': 'B', 'val_neta': 400, 'marja_bruta': 150, 'marja_neta': 150},
    ])
    with flask_app.app_context():
        ctx = overviews.build_context('team', AN, None, {})
    # 900 net, 350 margin -> 38.9%.
    assert ctx['cards'] == [
        ("Val. Netă", "900 RON"),
        ("Marjă Brută", "350 RON / 38.9%"),
        ("Marjă Netă", "350 RON / 38.9%"),
        ("Nr. Agenți", "2"),
    ]


def test_clients_context_shape(flask_app):
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('clients', AN, None, {})
    assert ctx['nav'] == 'clients'
    assert ctx['filename_base'] == f'clienti_{AN}'
    assert list(ctx['sheets']) == [f'Clienți {AN}']
    codes = [r['cod_client'] for r in _rows(ctx['sheets'][f'Clienți {AN}'])]
    assert 'C001' in codes and 'KAUFLAND' in codes


def test_clients_context_honours_brand_filter(flask_app):
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('clients', AN, None, {'brand': 'Basilur'})
    codes = [r['cod_client'] for r in _rows(ctx['sheets'][f'Clienți {AN}'])]
    assert codes == ['C001']


@pytest.mark.parametrize('report', ['team', 'clients'])
def test_context_honours_luna(flask_app, report):
    """Month 2 has no seeded rows, so its period sheet must come back empty
    while the unfiltered context has data. /export/team used to drop luna."""
    from exports import overviews
    with flask_app.app_context():
        full = overviews.build_context(report, AN, None, {})
        empty = overviews.build_context(report, AN, EMPTY_MONTH, {})
    first = list(full['sheets'])[0]
    assert _rows(full['sheets'][first])
    assert not _rows(empty['sheets'][first])
    assert empty['subtitle'] == '2026 · Februarie'


@pytest.mark.parametrize('report', ['team', 'clients'])
@pytest.mark.parametrize('luna', [13, -1])
def test_context_rejects_out_of_range_luna(flask_app, report, luna):
    from exports import overviews
    with flask_app.app_context():
        assert overviews.build_context(report, AN, luna, {}) is None


def test_context_unknown_report_returns_none(flask_app):
    from exports import overviews
    with flask_app.app_context():
        assert overviews.build_context('nope', AN, None, {}) is None


@pytest.mark.parametrize('report', ['team', 'clients'])
def test_context_luna_zero_is_year_to_date(flask_app, report):
    from exports import overviews
    with flask_app.app_context():
        zero = overviews.build_context(report, AN, 0, {})
        none = overviews.build_context(report, AN, None, {})
    assert zero['subtitle'] == none['subtitle']
    for name in none['sheets']:
        assert len(_rows(zero['sheets'][name])) == len(_rows(none['sheets'][name]))


def test_excel_sheet_is_not_truncated_to_the_deck_limit(flask_app, monkeypatch):
    """The deck caps at 15 rows and says so in its caption; the workbook must
    still carry every row."""
    import queries
    from exports import overviews
    fake = [{'client': f'Client {i}', 'cod_client': f'K{i}', 'agent': 'A',
             'tip_client': 'X', 'judet_client': 'Y', 'val_neta': 100 - i,
             'marja_bruta': 10, 'marja_bruta_pct': 10.0, 'marja_neta': 10,
             'marja_neta_pct': 10.0, 'ultima_comanda': '2026-01-15',
             'zile_inactiv': 1, 'nr_branduri': 1, 'delta_vn': None}
            for i in range(40)]
    monkeypatch.setattr(queries, 'clients_list', lambda *a, **k: fake)
    with flask_app.app_context():
        ctx = overviews.build_context('clients', AN, None, {})
    assert len(_rows(ctx['sheets'][f'Clienți {AN}'])) == 40
    assert len(ctx['tables'][0]['rows']) == 15
    assert ctx['tables'][0]['title'] == f'Clienți {AN} — top 15 din 40'


def test_products_context_shape(flask_app):
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('products', AN, None, {})
    assert ctx['nav'] == 'products'
    assert ctx['filename_base'] == f'produse_{AN}'
    assert list(ctx['sheets']) == ['Branduri', 'Top SKU']
    brands = [r['furnizor'] for r in _rows(ctx['sheets']['Branduri'])]
    assert 'Basilur' in brands and 'Toras' in brands
    assert 'SKU001' in [r['sku'] for r in _rows(ctx['sheets']['Top SKU'])]
    assert len(ctx['tables']) == 1
    assert len(ctx['extra_tables']) == 1
    assert ctx['extra_tables'][0]['title'].startswith(f'Top SKU {AN}')


def test_products_context_brand_filter_scopes_skus_only(flask_app):
    """The page filters only the SKU table by brand — the brand table always
    shows every brand. The export mirrors that."""
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('products', AN, None, {'brand': 'Basilur'})
    skus = _rows(ctx['sheets']['Top SKU'])
    assert skus
    assert {r['furnizor'] for r in skus} == {'Basilur'}
    assert 'Toras' in [r['furnizor'] for r in _rows(ctx['sheets']['Branduri'])]


def test_products_context_search_filter_reaches_the_sku_query(flask_app):
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('products', AN, None, {'q': 'SKU002'})
    assert [r['sku'] for r in _rows(ctx['sheets']['Top SKU'])] == ['SKU002']


def test_products_context_honours_luna(flask_app):
    from exports import overviews
    with flask_app.app_context():
        empty = overviews.build_context('products', AN, EMPTY_MONTH, {})
    assert not _rows(empty['sheets']['Branduri'])
    assert not _rows(empty['sheets']['Top SKU'])


def test_basilur_context_shape(flask_app):
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('basilur', AN, None, {})
    assert ctx['nav'] == 'basilur'
    assert ctx['filename_base'] == f'raportare_basilur_{AN}_YTD'
    assert list(ctx['sheets']) == [
        'Brand KPIs', 'Monthly Sales', 'Stock by Brand', 'Stock Detail']
    assert callable(ctx['ppt_builder'])
    assert ctx['curs'] == overviews.BASILUR_DEFAULT_CURS
    # The deck is bespoke, so the generic slide inputs stay empty.
    assert ctx['cards'] == []
    assert ctx['tables'] == []


def test_basilur_filename_names_the_month(flask_app):
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('basilur', AN, 1, {})
    assert ctx['filename_base'] == f'raportare_basilur_Ian_{AN}'


def test_basilur_curs_reaches_the_workbook(flask_app, monkeypatch):
    import queries
    from exports import overviews
    monkeypatch.setattr(queries, 'basilur_kpi_per_brand', lambda *a, **kw: [
        {'furnizor': 'Basilur', 'val_neta': 910.0, 'clienti_activi': 3,
         'nr_sku': 7, 'val_neta_py': 455.0, 'delta_vn': 100.0},
    ])
    with flask_app.app_context():
        ctx = overviews.build_context('basilur', AN, None, {'curs': 9.10})
    row = _rows(ctx['sheets']['Brand KPIs'])[0]
    assert row['Net Sales (USD)'] == 100     # 910 RON / 9.10
    assert row['Net Sales PY (USD)'] == 50   # 455 RON / 9.10
    assert ctx['curs'] == 9.10


@pytest.mark.parametrize('bad', [0, -1, None])
def test_basilur_falls_back_to_the_default_curs(flask_app, bad):
    """A zero or missing rate would divide by zero in every USD figure."""
    from exports import overviews
    with flask_app.app_context():
        ctx = overviews.build_context('basilur', AN, None, {'curs': bad})
    assert ctx['curs'] == overviews.BASILUR_DEFAULT_CURS


def test_build_basilur_ppt_uses_the_given_curs(flask_app):
    """Regression: the deck hard-coded 4.55 and ignored ?curs, so the workbook
    and the deck disagreed whenever the owner changed the rate."""
    from exports import ppt_export
    buf = ppt_export.build_basilur_ppt(
        an=AN, period_label='2026 YTD', kpi_total={'val_neta': 910.0},
        kpi_per_brand=[], monthly_data={}, stoc_per_brand=[], stoc_detail=[],
        curs=9.10)
    texts = [sh.text_frame.text for s in Presentation(buf).slides
             for sh in s.shapes if sh.has_text_frame]
    assert any('Rate: 1 USD = 9.1 RON' in t for t in texts)
    assert any('$100' in t for t in texts)  # 910 RON / 9.10


def test_build_basilur_ppt_lays_out_every_brand(flask_app):
    """Regression: the KPI-by-brand slide indexed a three-slot x-position list
    with four brands, so this builder raised IndexError on every call and
    /raportare-basilur/export/ppt returned a 500."""
    from exports import overviews, ppt_export
    buf = ppt_export.build_basilur_ppt(
        an=AN, period_label='2026 YTD', kpi_total={},
        kpi_per_brand=[{'furnizor': b, 'val_neta': 100.0, 'delta_vn': None,
                        'clienti_activi': 1, 'nr_sku': 1}
                       for b in overviews.BASILUR_BRANDS],
        monthly_data={}, stoc_per_brand=[], stoc_detail=[])
    texts = [sh.text_frame.text for s in Presentation(buf).slides
             for sh in s.shapes if sh.has_text_frame]
    for brand in overviews.BASILUR_BRANDS:
        assert brand.upper() in texts


REPORTS = ['team', 'clients', 'products', 'basilur']

EXPECTED_SHEETS = {
    'team': [f'Echipa {AN}', f'Echipa {AN - 1}'],
    'clients': [f'Clienți {AN}'],
    'products': ['Branduri', 'Top SKU'],
    'basilur': ['Brand KPIs', 'Monthly Sales', 'Stock by Brand', 'Stock Detail'],
}


@pytest.mark.parametrize('report', REPORTS)
def test_excel_route_returns_openable_workbook(client, report):
    rv = client.get(f'/export/{report}?an={AN}')
    assert rv.status_code == 200
    assert rv.mimetype == XLSX_MIME
    wb = openpyxl.load_workbook(io.BytesIO(rv.data))
    assert wb.sheetnames == EXPECTED_SHEETS[report]


@pytest.mark.parametrize('report', REPORTS)
def test_ppt_route_returns_openable_deck(client, report):
    rv = client.get(f'/export/ppt/{report}?an={AN}')
    assert rv.status_code == 200
    assert rv.mimetype == PPT_MIME
    assert len(Presentation(io.BytesIO(rv.data)).slides) >= 2


def test_team_excel_no_longer_ignores_luna(client):
    """Regression: /export/team called team_table(an) with no month filter, so
    it shipped a full year while the page showed a single month."""
    rv = client.get(f'/export/team?an={AN}&luna={EMPTY_MONTH}')
    assert rv.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(rv.data))
    assert wb[f'Echipa {AN}']['A1'].value == 'Nu există date pentru acest raport.'


def test_products_excel_honours_the_page_search_filter(client):
    """Regression: /export/products ignored ?q entirely."""
    rv = client.get(f'/export/products?an={AN}&q=SKU002')
    wb = openpyxl.load_workbook(io.BytesIO(rv.data))
    skus = [row[0] for row in wb['Top SKU'].iter_rows(min_row=2, values_only=True)]
    assert skus == ['SKU002']


@pytest.mark.parametrize('report', REPORTS)
@pytest.mark.parametrize('luna', [13, -1])
def test_routes_reject_out_of_range_luna(client, report, luna):
    assert client.get(f'/export/{report}?an={AN}&luna={luna}').status_code == 404
    assert client.get(f'/export/ppt/{report}?an={AN}&luna={luna}').status_code == 404


@pytest.mark.parametrize('report', REPORTS)
def test_routes_deny_role_without_nav_access(client, monkeypatch, report):
    import authz
    monkeypatch.setattr(authz, 'can_access_nav', lambda role, key: False)
    assert client.get(f'/export/{report}?an={AN}').status_code == 403
    assert client.get(f'/export/ppt/{report}?an={AN}').status_code == 403


def test_legacy_basilur_excel_url_redirects(client):
    rv = client.get(f'/raportare-basilur/export/excel?an={AN}&curs=5')
    assert rv.status_code == 302
    assert '/export/basilur' in rv.headers['Location']
    assert 'curs=5' in rv.headers['Location']


def test_legacy_basilur_ppt_url_redirects(client):
    rv = client.get(f'/raportare-basilur/export/ppt?an={AN}&curs=5')
    assert rv.status_code == 302
    assert '/export/ppt/basilur' in rv.headers['Location']
    assert 'curs=5' in rv.headers['Location']


def test_basilur_page_still_renders(client):
    """The page route now imports BASILUR_BRANDS from exports.overviews."""
    rv = client.get(f'/raportare-basilur?an={AN}')
    assert rv.status_code == 200
    assert 'KingsLeaf' in rv.get_data(as_text=True)
